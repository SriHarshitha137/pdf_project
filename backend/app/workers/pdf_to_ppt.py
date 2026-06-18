from pathlib import Path

from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches

from app.core.config import settings
from app.core.celery_app import celery_app
from app.db.session import SessionLocal
from app.workers.conversion_helpers import complete_job, fail_job, outputs_dir, start_file_job


@celery_app.task
def pdf_to_ppt_task(job_id: int):
    db = SessionLocal()
    job = None

    try:
        job, input_path = start_file_job(db, job_id)

        if not job:
            return

        dpi = int(job.options.get("dpi", 200))
        output_dir = outputs_dir() / f"pdf_to_ppt_{job_id}_images"
        output_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        images = convert_from_path(
            input_path,
            dpi=dpi,
            poppler_path=settings.POPPLER_PATH
        )

        if not images:
            raise Exception("PDF has no pages")

        presentation = Presentation()
        first_image = images[0]
        slide_width = Inches(10)
        slide_height = int(slide_width * first_image.height / first_image.width)
        presentation.slide_width = slide_width
        presentation.slide_height = slide_height

        blank_layout = presentation.slide_layouts[6]

        for index, image in enumerate(images, start=1):
            image_path = output_dir / f"page_{index}.png"
            image.save(image_path, "PNG")

            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height
            )

        output_path = outputs_dir() / f"pdf_to_ppt_{job_id}.pptx"
        presentation.save(output_path)

        complete_job(db, job, output_path)

    except Exception as e:
        fail_job(db, job, "PDF To PPT", job_id, e)

    finally:
        db.close()
